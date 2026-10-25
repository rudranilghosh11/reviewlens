from __future__ import annotations

from pathlib import Path
from io import BytesIO

from docx import Document
from fastapi import APIRouter, File, HTTPException, UploadFile
from fastapi.responses import PlainTextResponse

from backend.app.services.analysis import analyse_upload, search_reviews

router = APIRouter()
# No fixed application-level size cap by request. System RAM, disk and processing time still apply.
ALLOWED_TABLE_EXTENSIONS = {".csv", ".tsv", ".txt", ".xlsx", ".xlsm", ".xls", ".json", ".jsonl", ".ndjson", ".parquet", ".gz"}
_last_result: dict | None = None
_last_upload: tuple[bytes, str] | None = None

@router.get("/health")
def health():
    return {"status": "ok", "service": "review-lens-api", "upload_size_limit": None}

@router.post("/analyse")
async def analyse(file: UploadFile = File(...)):
    global _last_result, _last_upload
    filename = Path(file.filename or "upload.csv").name
    ext = Path(filename.lower()).suffix
    if not filename or ext not in ALLOWED_TABLE_EXTENSIONS:
        raise HTTPException(400, "Unsupported file type. Upload CSV, TSV, TXT, Excel, JSON/JSONL, Parquet, or a supported .gz file.")
    try:
        # UploadFile is spooled by Starlette; no fixed app-level byte cap is enforced here.
        file_bytes = await file.read()
        if not file_bytes:
            raise HTTPException(400, "The uploaded file is empty.")
        result = analyse_upload(file_bytes, filename)
        _last_upload = (file_bytes, filename)
        _last_result = result
        return result
    except HTTPException:
        raise
    except ImportError as exc:
        raise HTTPException(501, str(exc)) from exc
    except (ValueError, UnicodeError, OSError) as exc:
        raise HTTPException(400, str(exc)) from exc
    except Exception as exc:
        # Do not leak traceback or environment details to the client.
        print(f"Review Lens analysis failed for {filename}: {type(exc).__name__}: {exc}")
        raise HTTPException(500, "Analysis failed during ingestion or processing. Check the backend terminal for details.") from exc
    finally:
        await file.close()

@router.post("/analyse-sample")
def analyse_sample():
    global _last_result, _last_upload
    sample_path = Path(__file__).resolve().parents[3] / "data" / "reviews.csv"
    if not sample_path.is_file():
        raise HTTPException(404, "The bundled sample dataset is missing.")
    try:
        sample_bytes = sample_path.read_bytes()
        _last_upload = (sample_bytes, sample_path.name)
        _last_result = analyse_upload(sample_bytes, sample_path.name)
        return _last_result
    except Exception as exc:
        print(f"Sample analysis failed: {type(exc).__name__}: {exc}")
        raise HTTPException(500, "Sample analysis failed. Check the backend terminal for details.") from exc


@router.post("/product")
def analyse_product(product: str):
    """Re-run analysis for one product from the latest uploaded dataset."""
    global _last_result
    if _last_upload is None:
        raise HTTPException(404, "Upload and analyse a dataset first.")
    try:
        _last_result = analyse_upload(_last_upload[0], _last_upload[1], product_filter=product)
        return _last_result
    except (ValueError, UnicodeError, OSError) as exc:
        raise HTTPException(400, str(exc)) from exc
    except Exception as exc:
        print(f"Product analysis failed for {product}: {type(exc).__name__}: {exc}")
        raise HTTPException(500, "Product analysis failed. Check the backend terminal for details.") from exc

@router.get("/search")
def search_all_reviews(q: str, limit: int = 8):
    if not q.strip():
        raise HTTPException(400, "Enter a search query.")
    return {"query": q, "results": search_reviews(q, limit)}

@router.get("/results")
def latest_results():
    if _last_result is None:
        raise HTTPException(404, "No analysis has been run in this server session yet.")
    return _last_result

@router.get("/brief")
def download_brief():
    if _last_result is None:
        raise HTTPException(404, "No analysis brief is available yet.")
    return PlainTextResponse(
        _last_result["brief"],
        headers={"Content-Disposition": 'attachment; filename="reviewlens-brief.md"'},
    )


@router.get("/brief.docx")
def download_brief_docx():
    """Export the latest generated product-manager brief as a Word document."""
    if _last_result is None:
        raise HTTPException(404, "Run an analysis before downloading the brief.")

    document = Document()
    for line in _last_result["brief"].splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("# "):
            document.add_heading(stripped[2:], level=0)
        elif stripped.startswith("## "):
            document.add_heading(stripped[3:], level=1)
        elif stripped.startswith("### "):
            document.add_heading(stripped[4:], level=2)
        elif stripped.startswith("- "):
            document.add_paragraph(stripped[2:], style="List Bullet")
        elif len(stripped) > 2 and stripped[0].isdigit() and stripped[1:3] == ". ":
            document.add_paragraph(stripped[3:], style="List Number")
        else:
            # Strip lightweight Markdown emphasis markers for a clean Word document.
            document.add_paragraph(stripped.replace("**", "").replace("*", ""))

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    from fastapi.responses import StreamingResponse
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": 'attachment; filename="reviewlens-product-manager-brief.docx"'},
    )
